import sys
from pptx import Presentation
from pptx.util import Emu
p = Presentation(sys.argv[1])
print("slide size:", p.slide_width, "x", p.slide_height, "=", Emu(p.slide_width).pt, "x", Emu(p.slide_height).pt, "pt")
for i, s in enumerate(p.slides, 1):
    print(f"--- slide {i}: {len(s.shapes)} shapes ---")
    for sh in s.shapes:
        try:
            fill = sh.fill.type
            rgb = None
            if fill is not None and str(fill) == 'MSO_FILL_TYPE.SOLID (1)':
                rgb = sh.fill.fore_color.rgb
        except Exception as e:
            rgb = f"<{e}>"
        print(f"  id={sh.shape_id} name={sh.name!r} type={sh.shape_type} "
              f"L={Emu(sh.left).pt:.0f} T={Emu(sh.top).pt:.0f} "
              f"W={Emu(sh.width).pt:.0f} H={Emu(sh.height).pt:.0f} fill={rgb}")
